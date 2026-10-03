"""PowerPoint 簡報產生器（python-pptx，16:9）。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft JhengHei"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x57, 0x60, 0x6A)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
UP = RGBColor(0xC8, 0x42, 0x3B)
DOWN = RGBColor(0x2E, 0x8B, 0x57)
W, H = Inches(13.333), Inches(7.5)


def _run(p, text, size=14, bold=False, color=INK):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rpr = r._r.get_or_add_rPr()
    ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {})
        rpr.append(ea)
    ea.set("typeface", FONT)
    return r


class Deck:
    def __init__(self, footer: str):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.footer = footer
        self.n = 0

    def _blank(self, title: str | None = None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.n += 1
        if title:
            bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = NAVY
            bar.line.fill.background()
            tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
            _run(tb.text_frame.paragraphs[0], title, 26, True, NAVY)
        ft = s.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35))
        p = ft.text_frame.paragraphs[0]
        _run(p, f"{self.footer}　|　{self.n}", 10, False, MUTED)
        p.alignment = PP_ALIGN.RIGHT
        return s

    def text(self, s, x, y, w, h, items: list, size=14, title: str | None = None, fill=None):
        if fill is not None:
            box = s.shapes.add_shape(1, x, y, w, h)
            box.fill.solid()
            box.fill.fore_color.rgb = fill
            box.line.fill.background()
        tb = s.shapes.add_textbox(x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        # 依框面積與字數估算字級，避免溢出（每字約 size² 面積）
        chars = sum(len(str(i)) + 12 for i in items)
        area_pt2 = (w - Inches(0.3)) / 12700 * ((h - Inches(0.2)) / 12700) * 0.95
        while size > 9 and chars * (size * 1.25) ** 2 > area_pt2:
            size -= 0.5
        first = True
        if title:
            _run(tf.paragraphs[0], title, size + 2, True, NAVY)
            first = False
        for it in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(4)
            _run(p, "• " + it if not str(it).startswith(("•", "【")) else it, size, False, INK)
        return tb

    def title_slide(self, title: str, sub: str):
        s = self._blank()
        bg = s.shapes.add_shape(1, 0, 0, W, H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY
        bg.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.4))
        _run(tb.text_frame.paragraphs[0], title, 38, True, WHITE)
        tb2 = s.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.7), Inches(1))
        _run(tb2.text_frame.paragraphs[0], sub, 18, False, RGBColor(0xC9, 0xD6, 0xE8))

    def kpis(self, title: str, kpis: list[tuple[str, str, RGBColor | None]], notes: list[str]):
        s = self._blank(title)
        n = len(kpis)
        cw = (W - Inches(1.0) - Inches(0.25) * (n - 1)) / n
        for i, (label, value, color) in enumerate(kpis):
            x = Inches(0.5) + i * (cw + Inches(0.25))
            box = s.shapes.add_shape(1, x, Inches(1.4), cw, Inches(1.6))
            box.fill.solid()
            box.fill.fore_color.rgb = LIGHT
            box.line.fill.background()
            tb = s.shapes.add_textbox(x + Inches(0.15), Inches(1.5), cw - Inches(0.3), Inches(1.4))
            _run(tb.text_frame.paragraphs[0], label, 13, False, MUTED)
            _run(tb.text_frame.add_paragraph(), value, 30, True, color or INK)
        self.text(s, Inches(0.5), Inches(3.3), Inches(12.3), Inches(3.6), notes, 16)

    def image(self, title: str, img: Path | None, side: list[str] | None = None):
        if not img or not Path(img).exists():
            return
        s = self._blank(title)
        max_w = Inches(8.3) if side else Inches(12.3)
        max_h = Inches(5.7)
        pic = s.shapes.add_picture(str(img), Inches(0.4), Inches(1.2), height=max_h)
        if pic.width > max_w:  # 依長寬比縮到框內
            ratio = max_w / pic.width
            pic.width, pic.height = int(max_w), int(pic.height * ratio)
        if side:
            self.text(s, Inches(8.9), Inches(1.3), Inches(4.0), Inches(5.5), side, 13, fill=LIGHT)
        else:
            pic.left = int((W - pic.width) / 2)

    def table(self, title: str, header: list[str], rows: list[list], col_w: list[float] | None = None,
              color_col: int | None = None, size=12):
        s = self._blank(title)
        rows = rows[:14]
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.5), Inches(1.3), Inches(12.3),
                                   Inches(0.4) * (len(rows) + 1))
        t = shape.table
        if col_w:
            for i, w in enumerate(col_w):
                t.columns[i].width = Emu(int(Inches(12.3) * w / sum(col_w)))
        for i, h in enumerate(header):
            c = t.cell(0, i)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY
            c.text_frame.paragraphs[0].text = ""
            _run(c.text_frame.paragraphs[0], str(h), size, True, WHITE)
        for r, row in enumerate(rows, 1):
            for i, v in enumerate(row):
                c = t.cell(r, i)
                c.fill.solid()
                c.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
                c.text_frame.paragraphs[0].text = ""
                color = INK
                if color_col == i and isinstance(v, str) and v[:1] in "+-":
                    color = UP if v.startswith("+") else DOWN
                _run(c.text_frame.paragraphs[0], str(v), size, False, color)
        return s

    def order_slide(self, o: dict, note: str | None):
        qty = f"{o['lots']} 張" + (f" + {o['odd']} 股" if o["odd"] else "")
        if o["action"] == "BUY":
            s = self._blank(f"買進【{o.get('basket_name', '')}】{o['code']} {o['name']}｜{qty}｜限價 {o['limit_price']}")
            if note:
                self.text(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.7), [note], 13)
            self.text(s, Inches(0.5), Inches(1.9), Inches(6.0), Inches(2.9), o["why_stock"][:4], 12,
                      title="為何選擇這檔", fill=LIGHT)
            self.text(s, Inches(6.8), Inches(1.9), Inches(6.0), Inches(2.9), o["why_qty"], 12,
                      title="為何是這個股數", fill=LIGHT)
            mm = [f"【多買】{x}" for x in o["why_more"]] + [f"【少買】{x}" for x in o["why_less"]]
            self.text(s, Inches(0.5), Inches(4.95), Inches(6.0), Inches(2.0), mm, 11, title="多買／少買原因")
            self.text(s, Inches(6.8), Inches(4.95), Inches(6.0), Inches(2.0), o["exit_plan"][:3], 11,
                      title="出場計畫（賣出條件與張數）")
        else:
            s = self._blank(f"賣出 {o['code']} {o['name']}｜{qty}")
            self.text(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(2.4), [o["why_sell"]], 18,
                      title="賣出原因", fill=LIGHT)
            self.text(s, Inches(0.5), Inches(4.1), Inches(12.3), Inches(2.6),
                      [o["why_qty"], f"賣出後剩餘 {o['holding_after']:,} 股"], 16, title="為何賣這個股數", fill=LIGHT)

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.p.save(str(path))


def tuner_slide(deck: Deck, tune: dict, alpha_metrics: dict | None = None):
    st = tune["stats"]
    rows = [
        ["大盤隔夜跳空命中率", st["taiex_gap_n"], f"{st['taiex_gap_hit'] * 100:.1f}%" if st["taiex_gap_hit"] is not None else "累積中"],
        ["個股隔夜跳空命中率", st["stock_gap_n"], f"{st['stock_gap_hit'] * 100:.1f}%" if st["stock_gap_hit"] is not None else "累積中"],
        ["一個月推薦實際超額報酬", st["alpha_n"], f"{st['alpha_excess'] * 100:+.2f}%" if st["alpha_excess"] is not None else "需滿 20 交易日"],
        ["一個月推薦實際創新高比例", st["alpha_n"], f"{st['alpha_newhigh'] * 100:.1f}%" if st["alpha_newhigh"] is not None else "需滿 20 交易日"],
    ]
    if alpha_metrics and alpha_metrics.get("folds"):
        m = alpha_metrics
        rows.append(["回測：前 20 名創新高率 / 基準", len(m["folds"]), f"{m['topk_newhigh'] * 100:.1f}% / {m['base_rate'] * 100:.1f}%"])
        rows.append(["回測：前 20 名月超額報酬", len(m["folds"]), f"{m['topk_excess'] * 100:+.2f}%"])
    s = deck.table("模型健康度與調整 Agent", ["指標", "樣本", "表現"], rows, [3, 1, 2], size=12)
    notes = [f"[{i['severity']}] {i['detail']}" for i in tune["issues"]] + \
            [f"修正 {c['param']}：{c.get('reason', '')}" for c in tune["changes"]] + \
            ([tune["note"]] if tune.get("note") else ["目前未偵測到重大失誤，參數維持不變"])
    deck.text(s, Inches(0.5), Inches(4.6), Inches(12.3), Inches(2.3), notes, 12)


def build_daily(facts: dict, text: dict, charts: dict, path: Path) -> Path:
    deck = Deck(f"台股投資競賽 {facts['title']} {facts['date']}")
    deck.title_slide(f"{facts['title']}", f"{facts['date']}　|　2026 德明X富達投資擂台")
    pf = facts["portfolio"]
    tx = facts["predictions"].get("TAIEX", {})
    deck.kpis("今日重點", [
        ("總資產", f"{pf['nav'] / 1e4:,.0f} 萬", None),
        ("累積報酬", f"{pf['return_pct']:+.2f}%", UP if pf["return_pct"] > 0 else DOWN if pf["return_pct"] < 0 else INK),
        ("執行後持股率", f"{facts['projected_ratio'] * 100:.1f}%", None),
        ("大盤預測開盤", f"{tx.get('gap_pct', 0):+.2f}%", UP if tx.get("gap_pct", 0) > 0 else DOWN if tx.get("gap_pct", 0) < 0 else INK),
    ], [text["market_summary"], text["outlook"]])

    adv = facts.get("holding_advice") or []
    if adv:
        deck.table("庫存怎麼處理", ["標的", "持股", "損益", "建議", "數量", "限價"],
                   [[f"{h['code']} {h['name']}", f"{h['shares']:,}", f"{h['pnl_pct']:+.2f}%", h["verdict"], h["qty_text"],
                     h["limit_price"] or "—"] for h in adv], [2.2, 1.1, 1, 1, 4, 1], color_col=2, size=11)
    al_ = facts.get("alerts") or {}
    if al_.get("alerts"):
        lines = [f"【{a['level']}】{a['title']}：{a['reasons'][0]} → {a['action']}" for a in al_["alerts"][:8]]
        s = deck._blank(f"警示提醒（嚴重 {al_['counts'].get('嚴重', 0)}／警告 {al_['counts'].get('警告', 0)}）")
        deck.text(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.7), lines, 14)
    if al_.get("limit_up") or al_.get("limit_down"):
        rows = [[f"{'漲停' if q['chg_pct'] > 0 else '跌停'}｜{q['code']} {q['name']}", q.get("industry") or "", (q.get("reason") or "")[:60]]
                for q in (al_.get("limit_down", [])[:4] + al_.get("limit_up", [])[:8]) if q.get("reason")]
        if rows:
            deck.table(f"今日漲跌停與原因（漲停 {al_['limit_up_count']}／跌停 {al_['limit_down_count']}）", ["標的", "產業", "可能原因"],
                       rows, [2.2, 1.4, 6], size=11)

    al = facts["allocation"]
    etf_txt = f"／ETF {al['etf'] * 100:.0f}%" if al.get("etf") is not None else ""
    s = deck._blank(f"三籃配置：進攻 {al['offense'] * 100:.0f}%／穩健 {al['defense'] * 100:.0f}%{etf_txt}")
    deck.text(s, Inches(0.5), Inches(1.2), Inches(6.0), Inches(5.6), al["reasons"] + al.get("etf_reasons", []), 13,
              title="比例如何決定", fill=LIGHT)
    lines = []
    for b, name in (("offense", "進攻籃"), ("defense", "穩健籃"), ("etf", "ETF 籃")):
        v = al["baskets"].get(b)
        buys = [f"{o['code']} {o['name']}" for o in facts["orders"] if o["action"] == "BUY" and o.get("basket") == b]
        if v:
            lines.append(f"【{name}】目標 {v['target'] / 1e4:,.0f} 萬，執行後佔總資產 {v['after_pct']:.1f}%")
        lines.append(f"{name}買進：" + ("、".join(buys) if buys else "無"))
    lines.append(text.get("allocation_analysis", ""))
    deck.text(s, Inches(6.8), Inches(1.2), Inches(6.0), Inches(5.6), lines, 13, title="各籃內容")
    if facts.get("def_picks"):
        deck.table("穩健籃候選（收正報酬機率高、回檔小、不追高）", ["標的", "產業", "收正機率", "預期回檔", "β", "穩健分數", "訊號"],
                   [[f"{p['code']} {p['name']}", p["sector"], f"{p['p_up'] * 100:.1f}%", f"{p['exp_dd'] * 100:.1f}%",
                     f"{p['beta']:.2f}" if p["beta"] is not None else "-", f"{p['def_score']:+.2f}", p["signal"]]
                    for p in facts["def_picks"][:12]], [2.4, 1.6, 1, 1, 0.8, 1, 1], size=11)

    dv = facts.get("derivatives") or {}
    if dv:
        rows = [["做多" if o["action"] == "BUY" else "放空", o["contract"], str(o["qty"]), f"{o['limit_price']:,}", f"{o['stop']:,}",
                 f"{o['margin']:,}", f"{o['max_loss']:,}"] for o in dv.get("orders", [])]
        if dv.get("hedge"):
            h = dv["hedge"]
            rows.append(["買賣權避險", h["contract"], str(h["qty"]), str(h["limit_price"]), "-", "-", f"{h['premium']:,}"])
        if rows:
            deck.table(f"期貨與選擇權：{dv.get('summary')}", ["動作", "合約", "口數", "限價", "停損", "保證金", "最大損失"],
                       rows, [1.4, 2.2, 0.8, 1.2, 1.2, 1.4, 1.4], size=12)
        s2 = deck._blank(f"期貨與選擇權判斷：{dv.get('summary')}")
        why = [w for o in dv.get("orders", []) for w in o["why"]] + ((dv.get("hedge") or {}).get("why") or [])
        deck.text(s2, Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.6), (why or []) + (dv.get("reasons") or []), 13,
                  title="依據（模型估計，不是保證）", fill=LIGHT)

    att = facts["attribution"]
    title = f"為何今日會漲／會跌（{att['date']}）" + ("｜候選股示範" if facts["hypothetical"] else "")
    deck.image(title, charts.get("attribution"), [f"{r['code']} {r['name']}：{r['verdict']}" for r in att["rows"]][:6])
    if att["rows"]:
        s = deck.table("漲跌歸因明細", ["標的", "漲跌", "損益(元)", "大盤", "產業", "個股", "主因"],
                       [[f"{r['code']} {r['name']}", f"{r['ret_pct']:+.2f}%", f"{r['pnl']:+,}", f"{r['mkt_pct']:+.2f}%",
                         f"{r['ind_pct']:+.2f}%", f"{r['idio_pct']:+.2f}%", r["main_driver"]] for r in att["rows"]],
                       [2.4, 1, 1.4, 1, 1, 1, 1.6], color_col=1, size=11)
        deck.text(s, Inches(0.5), Inches(1.3 + 0.4 * (len(att["rows"]) + 1) + 0.2), Inches(12.3), Inches(2.2),
                  [text["daily_move_analysis"]], 12)
    deck.image("產業強弱", charts.get("sectors"))

    rec = facts["recovery"]
    s = deck._blank("虧損如何補回")
    deck.text(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.9), [text["recovery_plan"]], 13, fill=LIGHT)
    items = [f"【{it['decision']}】{it['code']} {it['name']}（{it['unrealized']:+,} 元）：{it['why']}" for it in rec["items"]]
    deck.text(s, Inches(0.5), Inches(3.3), Inches(12.3), Inches(3.6), rec["summary"] + (items or ["目前無持股"]), 12)

    s = deck._blank("國際大事對持股的影響")
    deck.text(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.7), [text["macro_analysis"]], 13, fill=LIGHT)
    ev_lines = [f"{e['event']} → {e['factor_name']}{'↑' if e['direction'] > 0 else '↓'}：預估 {e['total_amount']:+,} 元"
                f"（{e['total_pct_nav']:+.2f}%）" for e in facts["macro"]] or ["今日未辨識出重大國際事件"]
    deck.text(s, Inches(0.5), Inches(3.1), Inches(12.3), Inches(3.8), ev_lines, 12)
    for e in facts["macro"][:3]:
        if e["rows"]:
            deck.table(f"國際事件：{e['event'][:28]}", ["標的", "敏感度β", "預估影響", "金額(元)", "後續決策"],
                       [[f"{r['code']} {r['name']}", r["beta"], f"{r['impact_pct']:+.2f}%", f"{r['amount']:+,}", r["decision"]]
                        for r in e["rows"]], [2, 1, 1, 1.2, 5], color_col=2, size=11)

    deck.image("全市場一個月創新高潛力", charts.get("top_picks"),
               [f"掃描 {facts['alpha_model']['n_stocks']} 檔上市櫃股票與 ETF",
                f"全市場基準創新高率 {facts['alpha_model']['base_rate'] * 100:.1f}%",
                "模型：梯度提升樹，滾動時間外驗證"])
    deck.image("美股隔夜表現", charts.get("us"))
    deck.image("台股開盤預測（美股 → 台股）", charts.get("pred"),
               ["Ridge 迴歸：費半、那指期、台積電 ADR、VIX、匯率、產業對應美股",
                "命中率＝樣本外方向命中率", "低於 55% 或不如基準者降權"])
    rows = []
    for o in facts["orders"]:
        rows.append(["買進" if o["action"] == "BUY" else "賣出",
                     f"{'穩健' if o.get('basket') == 'defense' else '進攻'}｜{o['code']} {o['name']}",
                     f"{o['lots']} 張" + (f"+{o['odd']}股" if o["odd"] else ""), o["limit_price"],
                     o.get("stop", "-"), f"{o['amount']:,}", f"{o.get('weight_pct', 0):.1f}%" if o["action"] == "BUY" else "-"])
    if rows:
        deck.table("交易建議總表", ["動作", "標的", "數量", "限價", "停損", "金額", "權重"], rows,
                   [1, 2.4, 1.6, 1.2, 1.2, 1.8, 1])
    notes = {n["code"]: n["note"] for n in text["order_notes"]}
    for o in facts["orders"]:
        deck.order_slide(o, notes.get(o["code"]))
    deck.image("執行後持股配置", charts.get("allocation"),
               [f"期貨戶轉入證券戶：{facts['funding']['transfer_from_futures']:,} 元"] if facts["funding"]["transfer_from_futures"] else None)
    deck.table("競賽規則檢查", ["項目", "結果", "說明"], [[c["item"], c["status"], c["detail"]] for c in facts["compliance"]],
               [2.6, 0.9, 6], size=11)
    tuner_slide(deck, facts["tuner"], facts["alpha_model"].get("metrics"))
    deck.image("累積報酬走勢", charts.get("nav"))
    deck.save(path)
    return path
