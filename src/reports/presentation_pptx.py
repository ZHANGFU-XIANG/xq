"""決賽簡報投影片（16:9），沿用週報的版面系統（weekly_pptx：深藍封面、白底內頁、紅漲綠跌、原生圖表）。
每張投影片的講稿放在備忘稿；最後附 Q&A 預備頁（備用，不在 6 分鐘內）。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from src.reports.weekly_pptx import (DOWN, H, INK, INK2, LINE, MIST, MUTED, MX, UP, W, WHITE, _bullets, _rect, _table, _text,
                                     _title)

SOFT = RGBColor(0xB8, 0xC2, 0xD6)


def _pct(v, d=1):
    return "—" if v is None else f"{v * 100:+.{d}f}%"


def _tone(v):
    return INK if v is None else UP if v > 0 else DOWN if v < 0 else INK


def _slide(prs, title: str, sub: str | None, notes: str, page: int):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, Inches(0.12), INK)
    _title(s, title, sub, y=Inches(0.45))
    _text(s, W - MX - Inches(1), H - Inches(0.42), Inches(1), Inches(0.25), [(str(page), 9, MUTED, False, True)])
    _text(s, MX, H - Inches(0.42), Inches(8), Inches(0.25), [("2026 德明X富達投資擂台｜決賽簡報", 9, MUTED, False, False)])
    s.notes_slide.notes_text_frame.text = notes
    return s


def _stat(s, x, y, label, value, color=INK, w=Inches(2.8)):
    _text(s, x, y, w, Inches(0.3), [(label, 12, MUTED, False, False)])
    _text(s, x, y + Inches(0.3), w, Inches(0.6), [(value, 28, color, True, True)])


def build(f: dict, script: list[dict], qa: list[dict], path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    notes = {x["title"]: f"（約 {x['seconds']} 秒）{x['text']}" for x in script}
    page = 1
    # 封面
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, H, INK)
    _text(s, MX, Inches(1.3), Inches(10), Inches(0.4), [("2026 德明X富達投資擂台　全國大專院校組　決賽簡報", 14, SOFT, False, False)])
    _text(s, MX, Inches(1.8), Inches(11), Inches(1.2), [("多 Agent 資料驅動投資系統", 40, WHITE, True, False)])
    _text(s, MX, Inches(2.9), Inches(11), Inches(0.5), [("數字可追溯・風險可量化・錯誤能被系統發現", 18, SOFT, False, False)])
    _text(s, MX, Inches(3.9), Inches(6), Inches(1.4), [(f"{f['ret']:+.2f}%", 80, UP if f["ret"] > 0 else DOWN, True, True)])
    _text(s, MX, Inches(5.35), Inches(6), Inches(0.4), [(f"累積報酬（{f['start']}～{f['date']}）・總資產 {f['nav']:,} 元", 14, SOFT, False, False)])
    s.notes_slide.notes_text_frame.text = notes.get("封面", "")
    page += 1
    # 成果總覽：淨值 vs 加權
    s = _slide(prs, "成果總覽", "累積報酬與加權指數比較", notes.get("成果總覽", ""), page)
    page += 1
    _stat(s, MX, Inches(1.6), "累積報酬", f"{f['ret']:+.2f}%", _tone(f["ret"]))
    _stat(s, MX, Inches(2.6), "同期加權指數", "—" if f.get("taiex_ret") is None else f"{f['taiex_ret']:+.2f}%", _tone(f.get("taiex_ret")))
    _stat(s, MX, Inches(3.6), "成交筆數／標的", f"{f['trades']}／{f['symbols']}")
    _stat(s, MX, Inches(4.6), "持股率", "—" if f.get("holding_ratio") is None else f"{f['holding_ratio'] * 100:.1f}%")
    tx = {x["date"]: x["ret"] for x in f.get("taiex_series") or []}
    if f["nav_series"]:
        cd = CategoryChartData()
        cd.categories = [x["date"][5:] for x in f["nav_series"]]
        cd.add_series("本隊", [x["ret"] for x in f["nav_series"]])
        cd.add_series("加權指數", [tx.get(x["date"]) for x in f["nav_series"]])
        ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(3.8), Inches(1.5), Inches(8.9), Inches(5.2), cd).chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(12)
        from pptx.enum.chart import XL_MARKER_STYLE

        for i, col in enumerate((UP, MUTED)):
            ser = ch.series[i]
            ser.format.line.color.rgb = col
            ser.format.line.width = Pt(2.5 if i == 0 else 1.5)
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 6
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = col
            ser.marker.format.line.color.rgb = col
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = LINE
        ch.value_axis.tick_labels.font.size = Pt(10)
        ch.category_axis.tick_labels.font.size = Pt(10)
    # 策略
    p = f["params"]
    s = _slide(prs, "策略：全市場一個月選股＋結單轉股", "每天排名、集中持有、條件式換股", notes.get("策略", ""), page)
    page += 1
    _bullets(s, MX, Inches(1.6), Inches(12), Inches(4.5), [
        "全市場上市櫃股票與 ETF，每天以價量特徵預測未來 20 個交易日的超額報酬與創新高機率",
        f"集中持有綜合排名前 {p['max_positions']} 檔、等權配置；處置股、最新財報轉虧的股票不買",
        f"持有滿 {p['min_hold_days']} 個交易日後，出現扣成本後預期多 {p['swap_margin'] * 100:.0f}% 以上的股票就全部換股",
        f"獲利 ≥ {p['profit_swap_min_pnl'] * 100:.0f}% 且有更好標的時提前換股；最長持有 {p['expiry_days']} 個交易日",
        "盤後算好、隔天開盤前預掛：回測顯示開盤成交的報酬遠高於盤中追價"], size=18)
    # 驗證
    bt = f.get("backtest") or {}
    s = _slide(prs, "驗證：所有查得到的歷史、樣本外、實際做得到的成交", f"{bt.get('from', '')}～{bt.get('to', '')}，每年只用當時以前的資料重訓",
               notes.get("驗證", ""), page)
    page += 1
    _table(s, MX, Inches(1.7), Inches(12), ["情境", "年化報酬", "說明"], [
        ["隔日開盤成交（實際採用）", _pct(bt.get("open_cagr")), "盤後算預測、隔天開盤前預掛單"],
        ["同日收盤成交（舊假設）", _pct(bt.get("close_cagr")), "收盤算預測又用同一個收盤價成交，實戰做不到 → 已修正"],
        ["加權指數", _pct(bt.get("taiex_cagr")), "同期買進持有"]], [4, 2, 6], row_h=Inches(0.55),
        colors={(0, 1): _tone(bt.get("open_cagr")), (1, 1): MUTED}, num_cols={1})
    _bullets(s, MX, Inches(4.2), Inches(12), Inches(2), [
        f"任選 38 個交易日（一次競賽長度）報酬為正的比例 {(bt.get('win38') or 0) * 100:.0f}%；最大回檔 {_pct(bt.get('max_dd'), 0)}",
        "參數只在前半段挑選，後半段不能變差才採用；近 5 年模擬對參數過度敏感，不自動調參"], size=16)
    # 勝率
    cal = (bt.get("calibration") or {})
    if cal.get("buckets"):
        s = _slide(prs, "誠實的勝率", "每天依模型排名分組，隔日開盤買、20 天後賣、扣成本", notes.get("誠實的勝率", ""), page)
        page += 1
        _table(s, MX, Inches(1.7), Inches(12), ["當天排名", "樣本數", "賺錢機率", "平均報酬", "贏大盤機率"],
               [[f"第 {b['rank']} 名", f"{b['n']:,}", f"{b['win'] * 100:.1f}%", _pct(b["avg"], 2), f"{b['beat_market'] * 100:.1f}%"]
                for b in cal["buckets"]], [2, 2, 2, 2, 2], row_h=Inches(0.5), num_cols={1, 2, 3, 4})
        _text(s, MX, Inches(5.0), Inches(12), Inches(1.2), [(
            f"前 5 名分散持有賺錢機率 {cal.get('top5_portfolio_win', 0) * 100:.0f}%：優勢來自賺多賠少，所以分散、停損，不追求「每次都對」", 16, INK2, False, False)])
    # 架構
    s = _slide(prs, "系統架構：多 Agent 分工", "數字由程式計算，LLM 只負責分類與文字並經數字核對", notes.get("系統架構", ""), page)
    page += 1
    ag = f.get("agents") or []
    half = (len(ag) + 1) // 2
    for col, items in enumerate((ag[:half], ag[half:])):
        _bullets(s, MX + col * Inches(6.2), Inches(1.6), Inches(6), Inches(5),
                 [a["title"] + ("（LLM）" if a["llm"] else "") for a in items], size=13)
    # 風控
    s = _slide(prs, "風險控管", "每天壓力測試、追繳預警、黑天鵝指標回測", notes.get("風控", ""), page)
    page += 1
    if f.get("stress"):
        _table(s, MX, Inches(1.6), Inches(7.2), ["情境", "估計損益", "占總資產"],
               [[x["scenario"], f"{x['total']:+,}", f"{x['pct_nav']}%"] for x in f["stress"]], [3, 2, 1.4],
               row_h=Inches(0.45), colors={(i, 1): _tone(x["total"]) for i, x in enumerate(f["stress"])}, num_cols={1, 2})
    side = []
    if f.get("margin"):
        side.append(f"期貨距追繳 {f['margin']['pct_to_call']}%（台指 {f['margin']['call_level']:,} 點），即時推播")
    side += [f"{b['name']}：{b['label']}" for b in (f.get("swan") or [])[:4]]
    _bullets(s, Inches(8.1), Inches(1.6), Inches(4.7), Inches(5), side, size=13)
    # 交易檢討
    con = f.get("contrib") or []
    if con:
        s = _slide(prs, "交易檢討：損益貢獻", "已實現＋未實現，含手續費與證交稅", notes.get("交易檢討", ""), page)
        page += 1
        rows = con[:5] + ([None] if len(con) > 10 else []) + con[-5:] if len(con) > 10 else con
        _table(s, MX, Inches(1.6), Inches(9), ["股票", "損益（元）", "狀態"],
               [["…", "", ""] if r is None else [f"{r['code']} {r['name']}", f"{r['pnl']:+,}", "持有中" if r["held"] else "已出清"]
                for r in rows], [4, 2.5, 2], row_h=Inches(0.42),
               colors={(i, 1): _tone(r["pnl"]) for i, r in enumerate(rows) if r}, num_cols={1})
    # 影子組合
    sh = f.get("shadow") or {}
    if sh.get("breakdown"):
        s = _slide(prs, "影子組合：如果完全照策略", "同一套成本規則、開盤成交、不做期貨", notes.get("影子組合", ""), page)
        page += 1
        last = sh["last"]
        _stat(s, MX, Inches(1.6), "實際帳戶", f"{last['actual_ret']:+.2f}%", _tone(last["actual_ret"]))
        _stat(s, MX + Inches(3.2), Inches(1.6), "影子帳戶", f"{last['shadow_ret']:+.2f}%", _tone(last["shadow_ret"]))
        _stat(s, MX + Inches(6.4), Inches(1.6), "差距（元）", f"{sh['diff']:+,}", _tone(sh["diff"]))
        _table(s, MX, Inches(3.0), Inches(9), ["差距來源", "金額（元）"], [[b["name"], f"{b['amount']:+,}"] for b in sh["breakdown"]],
               [5, 2], row_h=Inches(0.48), colors={(i, 1): _tone(b["amount"]) for i, b in enumerate(sh["breakdown"])}, num_cols={1})
    # 發現與修正
    s = _slide(prs, "發現問題、驗證、修正", "系統自己會抓錯", notes.get("發現與修正", ""), page)
    page += 1
    _table(s, MX, Inches(1.6), Inches(12), ["發現", "原因", "修正"], [
        ["盤後報告用到前一天價格", "雲端沿用盤中快取，未滿 3 小時不更新", "收盤後缺當天資料就強制更新＋資料檢查 Agent"],
        ["回測年化被高估", "假設收盤價成交，實戰做不到", "改隔日開盤成交、全歷史樣本外驗證"],
        ["持有現金與官方差 1 元", "交割金額賣出無條件進位、損益四捨五入", "兩種口徑分開計算，對帳完全一致"],
        ["網站重啟要 2 分鐘", "每次重抓 10 年資料", "雲端打包快照，3 秒啟動"]], [3.5, 4, 4.5], row_h=Inches(0.6))
    # 結論
    s = _slide(prs, "結論", None, notes.get("結論", ""), page)
    page += 1
    _bullets(s, MX, Inches(1.8), Inches(12), Inches(4), [
        "決策流程可驗證：每個數字都能追溯到程式與官方資料",
        "風險可量化：壓力測試、追繳點位、影子組合把改進方向換算成金額",
        "誠實面對不確定：公開真實勝率，不追求不存在的高準確率",
        "下一步：降低期貨槓桿、開盤前預掛減少滑價、集中持有前 5 名"], size=20)
    # Q&A 備用
    s = _slide(prs, "Q&A 預備（備用）", "答案中的數字皆由系統產生", "", page)
    _table(s, MX, Inches(1.4), Inches(12), ["問題", "回答重點"], [[x["q"], x["a"]] for x in qa], [3, 9], row_h=Inches(0.5))
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
