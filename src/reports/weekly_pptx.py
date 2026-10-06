"""投資競賽週報（16:9 .pptx）。

版面系統（frontend-design 規劃）：
- 主角是每頁頂端的「週帶」：週一到週五五格，每格是當天加權指數漲跌（台股紅漲綠跌）；講某筆交易時，
  那一天的格子會被框起來，讀者一眼知道交易發生在哪天、當天盤勢如何。
- 顏色沿用網站：進攻紫、穩健藍、ETF 青、期貨琥珀、上漲紅、下跌綠；內頁白底，封面與章節頁用深藍。
- 字型：中文微軟正黑體，數字 Bahnschrift（窄體、數字好讀）；一律靠左對齊。
- 圖表用 PowerPoint 原生圖表（下載後仍可編輯）。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x14, 0x21, 0x3D)
INK2 = RGBColor(0x3B, 0x47, 0x63)
MUTED = RGBColor(0x6B, 0x74, 0x89)
LINE = RGBColor(0xD9, 0xDE, 0xE7)
MIST = RGBColor(0xEE, 0xF1, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
UP = RGBColor(0xC8, 0x32, 0x3A)
DOWN = RGBColor(0x17, 0x8A, 0x55)
UP_ON_DARK = RGBColor(0xFF, 0x7A, 0x80)
DOWN_ON_DARK = RGBColor(0x5F, 0xD9, 0x9C)
ASSET = {"STOCK": RGBColor(0x6B, 0x46, 0xA8), "ETF": RGBColor(0x2E, 0x7D, 0x6B), "FUT": RGBColor(0xB4, 0x56, 0x0F),
         "OPT": RGBColor(0xB4, 0x56, 0x0F)}
ASSET_NAME = {"STOCK": "股票", "ETF": "ETF", "FUT": "期貨", "OPT": "選擇權"}
CJK = "Microsoft JhengHei"
NUM = "Bahnschrift"
W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)


def _font(run, size: float, color: RGBColor = INK, bold: bool = False, num: bool = False) -> None:
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = NUM if num else CJK
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", NUM if (num and tag == "a:latin") else CJK)


def _text(slide, x, y, w, h, parts, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing: float = 1.15):
    """parts：[(文字, 大小, 顏色, 粗體, 數字字型)] 或多段 [[...], [...]]（每段一個段落）。"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = parts if parts and isinstance(parts[0], list) else [parts]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        for seg in para:
            txt, size, color, bold, num = (list(seg) + [INK, False, False])[:5] if len(seg) < 5 else seg
            r = p.add_run()
            r.text = txt
            _font(r, size, color, bold, num)
    return tb


def _rect(slide, x, y, w, h, fill: RGBColor | None, line: RGBColor | None = None, width: float = 0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def _num(v, digits: int | None = None) -> str:
    """價格與金額：千分位；整數不顯示 .0，小數最多 2 位。"""
    if v is None:
        return "—"
    v = float(v)
    if digits is None:
        digits = 0 if v == int(v) else 2
    return f"{v:,.{digits}f}"


def _pct(v, digits: int = 2) -> str:
    return "—" if v is None else f"{v:+.{digits}f}%"


def _tone(v, dark: bool = False) -> RGBColor:
    if v is None or v == 0:
        return WHITE if dark else INK2
    if dark:
        return UP_ON_DARK if v > 0 else DOWN_ON_DARK
    return UP if v > 0 else DOWN


def _mmdd(iso: str) -> str:
    return f"{int(iso[5:7])}/{int(iso[8:10])}"


def week_tape(slide, daily: list[dict], highlight: str | None = None, dark: bool = False, y=Inches(0.32)) -> None:
    gap = Inches(0.1)
    cw = int((W - 2 * MX - gap * 4) / 5)
    for i, d in enumerate(daily):
        x = MX + i * (cw + gap)
        on = highlight == d["date"]
        bg = (RGBColor(0x22, 0x33, 0x55) if dark else MIST)
        box = _rect(slide, x, y, cw, Inches(0.62), bg, (WHITE if dark else INK) if on else None, 2.25 if on else 0.75,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        box.adjustments[0] = 0.12
        label_color = (RGBColor(0xB8, 0xC2, 0xD6) if dark else MUTED)
        _text(slide, x + Inches(0.16), y + Inches(0.08), Inches(1.2), Inches(0.24),
              [(f"週{d['weekday']}　", 10, label_color, False, False), (_mmdd(d["date"]), 10, label_color, False, True)])
        if d.get("trading") and d.get("taiex_chg_pct") is not None:
            val = [(_pct(d["taiex_chg_pct"]), 17, _tone(d["taiex_chg_pct"], dark), True, True)]
        elif d.get("trading"):
            val = [("盤中", 13, label_color, False, False)]
        else:
            val = [("休市", 13, label_color, False, False)]
        _text(slide, x + Inches(0.16), y + Inches(0.27), cw - Inches(0.3), Inches(0.32), val)
        _text(slide, x + cw - Inches(1.05), y + Inches(0.08), Inches(0.9), Inches(0.24),
              [("加權", 9, label_color, False, False)], align=PP_ALIGN.RIGHT)


def _title(slide, text: str, sub: str | None = None, y=Inches(1.18)) -> None:
    _text(slide, MX, y, W - 2 * MX, Inches(0.6), [(text, 26, INK, True, False)])
    if sub:
        _text(slide, MX, y + Inches(0.62), W - 2 * MX, Inches(0.35), [(sub, 13, MUTED, False, False)])


def _footer(slide, w: dict, page: int) -> None:
    _text(slide, MX, H - Inches(0.42), Inches(8), Inches(0.25),
          [(f"2026 德明X富達投資擂台｜{w['week_label']}週報｜{w['monday']}～{w['end_day']}", 9, MUTED, False, False)])
    _text(slide, W - MX - Inches(1), H - Inches(0.42), Inches(1), Inches(0.25), [(str(page), 9, MUTED, False, True)],
          align=PP_ALIGN.RIGHT)


def _bullets(slide, x, y, w, h, items: list[str], size: float = 14, color: RGBColor = INK2, marker: RGBColor = INK) -> None:
    paras = [[("■ ", size * 0.55, marker, False, False), (it, size, color, False, False)] for it in items if it]
    if paras:
        tb = _text(slide, x, y, w, h, paras, line_spacing=1.3)
        for p in tb.text_frame.paragraphs:
            p.space_after = Pt(size * 0.55)


def _table(slide, x, y, w, header: list[str], rows: list[list], widths: list[float], *, row_h=Inches(0.36),
           colors: dict | None = None, num_cols: set[int] = frozenset()) -> None:
    shape = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, row_h * (len(rows) + 1))
    t = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set("firstRow", "0")
    tbl_pr.set("bandRow", "0")
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 無樣式、無格線（只畫自訂的列底線）
    total = sum(widths)
    for i, wd in enumerate(widths):
        t.columns[i].width = Emu(int(w * wd / total))
    for r_i in range(len(rows) + 1):
        t.rows[r_i].height = row_h
        for c_i in range(len(header)):
            cell = t.cell(r_i, c_i)
            cell.fill.background() if r_i else cell.fill.solid()
            if r_i == 0:
                cell.fill.fore_color.rgb = MIST
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            val = header[c_i] if r_i == 0 else rows[r_i - 1][c_i]
            r = tf.paragraphs[0].add_run()
            r.text = "" if val is None else str(val)
            color = MUTED if r_i == 0 else (colors or {}).get((r_i - 1, c_i), INK)
            _font(r, 11 if r_i == 0 else 12, color, False, c_i in num_cols and r_i > 0)
            if c_i in num_cols:
                tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
            # 只留每列底線（其他三邊不畫）
            tc_pr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tc_pr.makeelement(qn(tag), {"w": "6350" if tag == "a:lnB" else "0"})
                if tag == "a:lnB":
                    fill = ln.makeelement(qn("a:solidFill"), {})
                    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": "D9DEE7"}))
                else:
                    fill = ln.makeelement(qn("a:noFill"), {})
                ln.append(fill)
                tc_pr.append(ln)


# ---------------------------------------------------------------- 各頁
def _cover(prs, w: dict, text) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, H, INK)
    _text(s, MX, Inches(1.2), Inches(8), Inches(0.4), [("2026 德明X富達投資擂台　全國大專院校組", 14, RGBColor(0xB8, 0xC2, 0xD6), False, False)])
    _text(s, MX, Inches(1.65), Inches(9), Inches(0.9),
          [(f"{w['week_label']}週報", 40, WHITE, True, False)])
    _text(s, MX, Inches(2.5), Inches(9), Inches(0.4),
          [(f"{w['monday']}～{w['end_day']}", 16, RGBColor(0xB8, 0xC2, 0xD6), False, True)])
    _text(s, MX, Inches(3.2), Inches(6.5), Inches(1.5), [(_pct(w["week_ret_pct"]), 88, _tone(w["week_ret_pct"], True), True, True)])
    _text(s, MX, Inches(4.75), Inches(6.5), Inches(0.35), [("本週總資產報酬", 14, RGBColor(0xB8, 0xC2, 0xD6), False, False)])
    stats = [("加權指數", _pct(w["taiex_week_pct"]), _tone(w["taiex_week_pct"], True)),
             ("累積報酬", _pct(w["cum_ret_pct"]), _tone(w["cum_ret_pct"], True)),
             ("期末總資產", f"{w['nav_end']:,}", WHITE),
             ("持股率", f"{w['holding_ratio'] * 100:.1f}%", WHITE if w["holding_ratio"] >= 0.7 else UP_ON_DARK),
             ("本週成交", f"{w['trade_count_week']} 筆", WHITE)]
    for i, (k, v, c) in enumerate(stats):
        x = Inches(7.4) + (i % 2) * Inches(2.8)
        y = Inches(3.25) + (i // 2) * Inches(0.95)
        _text(s, x, y, Inches(2.7), Inches(0.3), [(k, 12, RGBColor(0xB8, 0xC2, 0xD6), False, False)])
        _text(s, x, y + Inches(0.3), Inches(2.7), Inches(0.5), [(v, 24, c, True, True)])
    _text(s, MX, Inches(5.55), W - 2 * MX, Inches(0.5), [(text.headline, 16, WHITE, False, False)])
    week_tape(s, w["daily"], dark=True, y=Inches(6.35))


def _overview(prs, w: dict, text, page: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    week_tape(s, w["daily"])
    _title(s, "本週重點", text.headline)
    _bullets(s, MX, Inches(2.3), Inches(7.2), Inches(3.8), text.summary, size=16)
    x0 = Inches(8.4)
    _rect(s, x0 - Inches(0.3), Inches(2.25), Pt(1.5), Inches(4.2), LINE)
    kpis = [("本週總資產", _pct(w["week_ret_pct"]), _tone(w["week_ret_pct"])),
            ("同期加權指數", _pct(w["taiex_week_pct"]), _tone(w["taiex_week_pct"])),
            ("累積報酬（排名只看正報酬）", _pct(w["cum_ret_pct"]), _tone(w["cum_ret_pct"])),
            ("持股率（每日收盤須 ≥ 70%）", f"{w['holding_ratio'] * 100:.1f}%", INK if w["holding_ratio"] >= 0.7 else UP),
            ("累計成交／門檻 10 筆", f"{w['trade_count_total']} 筆", INK if w["trade_count_total"] >= 10 else UP),
            ("交易過的標的／門檻 5 檔", f"{w['symbols_total']} 檔", INK if w["symbols_total"] >= 5 else UP)]
    for i, (k, v, c) in enumerate(kpis):
        y = Inches(2.25) + i * Inches(0.7)
        _text(s, x0, y, Inches(3), Inches(0.25), [(k, 11, MUTED, False, False)])
        _text(s, x0, y + Inches(0.24), Inches(4.2), Inches(0.4), [(v, 20, c, True, True)])
    _footer(s, w, page)


def _market(prs, w: dict, text, page: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    week_tape(s, w["daily"])
    _title(s, "每日市場與總資產", "總資產與加權指數都以上週五收盤為 100")
    days = [d for d in w["daily"] if d["trading"] and d["nav"] is not None]
    if days:
        cd = CategoryChartData()
        cd.categories = [f"週{d['weekday']}" for d in days]
        cd.add_series("總資產", [round(d["nav"] / w["nav_start"] * 100, 2) for d in days])
        cd.add_series("加權指數", [d["taiex_idx"] for d in days])
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, MX, Inches(2.25), Inches(6.2), Inches(4.3), cd)
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
        for ser, color in zip(ch.series, (ASSET["STOCK"], MUTED)):
            ser.format.line.color.rgb = color
            ser.format.line.width = Pt(2.25)
            ser.smooth = False
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = color
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = LINE
        ch.value_axis.tick_labels.font.size = Pt(10)
        ch.category_axis.tick_labels.font.size = Pt(11)
        ch.value_axis.format.line.fill.background()
    _bullets(s, Inches(7.2), Inches(2.3), Inches(5.5), Inches(4.4), text.market_review, size=13)
    _footer(s, w, page)


def _trade_table(prs, w: dict, page: int) -> int:
    trades = w["trades"]
    chunks = [trades[i:i + 10] for i in range(0, len(trades), 10)] or [[]]
    for n, chunk in enumerate(chunks):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        week_tape(s, w["daily"])
        _title(s, "本週交易明細" + (f"（{n + 1}/{len(chunks)}）" if len(chunks) > 1 else ""),
               f"共 {len(trades)} 筆；理由來源標示系統建議、盤中建議、記錄備註或自行判斷")
        if not chunk:
            _text(s, MX, Inches(2.6), Inches(10), Inches(1), [("本週沒有成交紀錄。在 XQ 成交後，用網站「持股」頁匯入或跟 AI 助理說成交內容，下週週報就會出現。", 16, INK2, False, False)])
        else:
            rows, colors = [], {}
            for i, t in enumerate(chunk):
                rows.append([f"週{t['weekday']} {_mmdd(t['date'])}", "買進" if t["side"] == "BUY" else "賣出", ASSET_NAME[t["asset"]],
                             f"{t['code']} {t['name']}" if t["asset"] in ("STOCK", "ETF") else t["name"],
                             f"{t['qty']:,} {t['unit']}", _num(t["price"]), _num(t["amount"], 0), t["reason"]["source"]])
                colors[(i, 1)] = UP if t["side"] == "BUY" else DOWN
                colors[(i, 2)] = ASSET[t["asset"]]
            _table(s, MX, Inches(2.3), W - 2 * MX, ["日期", "動作", "類別", "標的", "數量", "價格", "金額（元）", "理由來源"], rows,
                   [1.2, 0.7, 0.8, 2.2, 1.2, 1.1, 1.4, 2.6], colors=colors, num_cols={4, 5, 6})
        _footer(s, w, page)
        page += 1
    return page


def _trade_details(prs, w: dict, text, page: int) -> int:
    notes = {(n.code, n.date): n.comment for n in text.trade_notes}
    for t in w["trades"]:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        week_tape(s, w["daily"], highlight=t["date"])
        verb = "買進" if t["side"] == "BUY" else "賣出"
        name = f"{t['code']} {t['name']}" if t["asset"] in ("STOCK", "ETF") else t["name"]
        _rect(s, MX, Inches(1.25), Pt(5), Inches(0.95), ASSET[t["asset"]])
        _text(s, MX + Inches(0.22), Inches(1.18), Inches(11), Inches(0.6),
              [(f"週{t['weekday']}　{verb} {name}", 26, INK, True, False)])
        _text(s, MX + Inches(0.22), Inches(1.8), Inches(11), Inches(0.4),
              [(f"{ASSET_NAME[t['asset']]}　", 13, ASSET[t["asset"]], True, False),
               (f"{t['qty']:,} {t['unit']} × {_num(t['price'])} ＝ {_num(t['amount'], 0)} 元", 13, INK2, False, True)])
        r = t["reason"]
        comment = notes.get((t["code"], t["date"])) or "；".join(r.get("reasons", [])[:1]) or r["source"]
        _text(s, MX, Inches(2.55), Inches(7.4), Inches(0.3), [("為什麼這樣做", 12, MUTED, False, False)])
        details = [x for x in r.get("reasons", [])[:3] if x and x not in comment]
        details += [f"出場計畫：{x}" for x in r.get("exit_plan", [])[:1]] + [f"風險：{x}" for x in r.get("risks", [])[:1]]
        paras = [[(comment, 17, INK, False, False)]]
        paras += [[("■ ", 7, INK, False, False), (d, 12, INK2, False, False)] for d in details]
        paras += [[(f"理由來源：{r['source']}", 11, MUTED, False, False)]]
        if r["source"].startswith("不在系統建議"):
            paras += [[("下次記錄成交時一併說明理由（例如跟 AI 助理說「因為法說會釋出好消息，買進 2330 一張」），週報就會寫進來。", 11, MUTED, False, False)]]
        tb = _text(s, MX, Inches(2.88), Inches(7.4), Inches(3.6), paras, line_spacing=1.3)
        for i, para in enumerate(tb.text_frame.paragraphs):
            para.space_after = Pt(14 if i == 0 else 6)
        x1 = Inches(8.6)
        _rect(s, x1 - Inches(0.3), Inches(2.55), Pt(1.5), Inches(3.9), LINE)
        m = t["market"]
        _text(s, x1, Inches(2.55), Inches(4), Inches(0.3), [("當天市場", 12, MUTED, False, False)])
        _text(s, x1, Inches(2.85), Inches(4), Inches(0.5),
              [("加權 ", 14, INK2, False, False), (_pct(m["taiex_chg_pct"]), 22, _tone(m["taiex_chg_pct"]), True, True)])
        lines = [e["event"] for e in m["events"][:2]] + [x["event"] for x in m["macro"][:1]]
        _bullets(s, x1, Inches(3.5), Inches(4.1), Inches(1.6), lines or ["當天報告沒有記錄重大事件"], size=11)
        if r.get("plan_price"):
            diff = (t["price"] / r["plan_price"] - 1) * 100
            _text(s, x1, Inches(5.25), Inches(4.1), Inches(0.9),
                  [[("系統建議價 ", 11, MUTED, False, False), (_num(r["plan_price"]), 13, INK, True, True),
                    ("　實際成交 ", 11, MUTED, False, False), (_num(t["price"]), 13, INK, True, True)],
                   [("差距 ", 11, MUTED, False, False), (_pct(diff), 13, _tone(-diff if t["side"] == "BUY" else diff), True, True),
                    ("（買進低於建議價為佳）" if t["side"] == "BUY" else "（賣出高於建議價為佳）", 10, MUTED, False, False)]])
        _footer(s, w, page)
        page += 1
    return page


def _holdings(prs, w: dict, page: int) -> int:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    week_tape(s, w["daily"])
    _title(s, "期末持股", f"總資產 {w['nav_end']:,} 元；證券戶現金 {w['cash']:,} 元；期貨戶權益 {w['futures_equity']:,} 元")
    hs = w["holdings"][:11]
    if not hs:
        _text(s, MX, Inches(2.6), Inches(10), Inches(1), [("期末沒有持股。", 16, INK2, False, False)])
    else:
        rows, colors = [], {}
        for i, h in enumerate(hs):
            rows.append([ASSET_NAME.get(h["asset"], h["asset"]), f"{h['code']} {h['name']}" if h["asset"] in ("STOCK", "ETF") else h["name"],
                         f"{h['shares']:,}", _num(h["avg_cost"]), _num(h["price"]),
                         _num(h["pnl"], 0), _pct(h["pnl_pct"]) if h["pnl_pct"] is not None else "—",
                         f"{h['weight_pct']:.1f}%" if h["weight_pct"] is not None else "—"])
            colors[(i, 0)] = ASSET.get(h["asset"], INK)
            colors[(i, 5)] = _tone(h["pnl"])
            colors[(i, 6)] = _tone(h["pnl_pct"]) if h["pnl_pct"] is not None else MUTED
        _table(s, MX, Inches(2.3), W - 2 * MX, ["類別", "標的", "股數／口數", "成本", "現價", "未實現損益", "報酬", "權重"], rows,
               [0.8, 2.6, 1.2, 1.1, 1.1, 1.5, 1.0, 0.9], colors=colors, num_cols={2, 3, 4, 5, 6, 7})
    _footer(s, w, page)
    return page + 1


def _contrib(prs, w: dict, page: int) -> int:
    rows = w["contrib"]
    if not rows:
        return page
    s = prs.slides.add_slide(prs.slide_layouts[6])
    week_tape(s, w["daily"])
    _title(s, "本週損益來源", "各檔本週損益（含已賣出部分）；紅色賺、綠色賠")
    rows = list(reversed(rows[:8]))
    cd = CategoryChartData()
    cd.categories = [f"{r['code']} {r['name']}" for r in rows]
    cd.add_series("本週損益（元）", [r["pnl"] for r in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MX, Inches(2.2), W - 2 * MX, Inches(4.5), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.number_format = '#,##0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(11)
    for i, r in enumerate(rows):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = UP if r["pnl"] > 0 else DOWN
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = LINE
    ch.value_axis.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.size = Pt(11)
    _footer(s, w, page)
    return page + 1


def _plan(prs, w: dict, text, page: int) -> int:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    week_tape(s, w["daily"])
    _title(s, "檢討與下週計畫")
    _text(s, MX, Inches(2.0), Inches(5.6), Inches(0.35), [("本週檢討", 15, INK, True, False)])
    _bullets(s, MX, Inches(2.5), Inches(5.6), Inches(4), text.lessons, size=14)
    x1 = Inches(6.9)
    _rect(s, x1 - Inches(0.3), Inches(2.0), Pt(1.5), Inches(4.5), LINE)
    _text(s, x1, Inches(2.0), Inches(5.8), Inches(0.35), [("下週計畫", 15, INK, True, False)])
    _bullets(s, x1, Inches(2.5), Inches(5.8), Inches(4), text.next_week, size=14)
    _footer(s, w, page)
    return page + 1


def _notes(prs, w: dict, verified: bool, page: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, "資料來源與說明", y=Inches(0.8))
    _bullets(s, MX, Inches(1.8), W - 2 * MX, Inches(5), [
        "成交：XQ 全球贏家成交回報（網站匯入或 AI 助理記錄）；同一筆委託分多次成交時合併為一列，價格為成交量加權平均",
        "總資產採競賽 App「投資組合總值」口徑：初始 1,000 萬＋未實現＋已實現；股票用當天收盤價、期貨用該合約日盤結算價（與 App 對帳一致）",
        "市場資料：證交所／櫃買中心日 K、加權指數、期交所行情；事件與國際大事來自當天盤前報告收錄的新聞",
        "交易理由：優先引用當天系統建議的理由；沒有對應建議者標示「自行判斷」，不另行補寫",
        ("文字說明由開源 AI 模型依上述資料撰寫，所有數字已逐一與資料核對" if verified
         else "文字說明採固定格式自動產生（AI 服務暫時無法使用或未通過數字核對）"),
        "持股率同 App：股票＋ETF 市值 ÷（市值＋持有現金＋期貨保證金＋期貨未實現）；損益同 App 不扣手續費，手續費另計",
    ], size=14)
    _footer(s, w, page)


def build(w: dict, text, verified: bool, path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    _cover(prs, w, text)
    page = 2
    _overview(prs, w, text, page)
    page += 1
    _market(prs, w, text, page)
    page += 1
    page = _trade_table(prs, w, page)
    page = _trade_details(prs, w, text, page)
    page = _holdings(prs, w, page)
    page = _contrib(prs, w, page)
    page = _plan(prs, w, text, page)
    _notes(prs, w, verified, page)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path
