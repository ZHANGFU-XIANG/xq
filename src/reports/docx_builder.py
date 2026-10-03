"""Word 報告產生器（python-docx）。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.config import settings

FONT = "Microsoft JhengHei"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
UP = RGBColor(0xC8, 0x42, 0x3B)
DOWN = RGBColor(0x2E, 0x8B, 0x57)


def _font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, hex_fill: str):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)
        style = self.d.styles["Normal"]
        style.font.name = FONT
        style.font.size = Pt(10.5)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
            st = self.d.styles[name]
            st.font.name = FONT
            st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
            st.font.color.rgb = NAVY

    def title(self, text: str, sub: str):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(text), 20, True, NAVY)
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(sub), 10, False, RGBColor(0x57, 0x60, 0x6A))

    def h1(self, text):
        self.d.add_heading(text, level=1)

    def h2(self, text):
        self.d.add_heading(text, level=2)

    def para(self, text, bold=False, color=None, size=None):
        p = self.d.add_paragraph()
        _font(p.add_run(text), size, bold, color)
        return p

    def bullets(self, items, label: str | None = None):
        if label:
            self.para(label, bold=True)
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            _font(p.add_run(it))

    def link_bullet(self, text, url):
        p = self.d.add_paragraph(style="List Bullet")
        _font(p.add_run(text + "　"))
        part = self.d.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "1F5FAD")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(c)
        rpr.append(u)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = "[來源]"
        r.append(t)
        h.append(r)
        p._p.append(h)

    def table(self, header: list[str], rows: list[list], widths: list[float] | None = None, color_col: int | None = None):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            _font(cell.paragraphs[0].add_run(str(h)), 9.5, True, RGBColor(0xFF, 0xFF, 0xFF))
            _shade(cell, "1F3A5F")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                color = None
                if color_col is not None and i == color_col and isinstance(v, str) and v[:1] in "+-":
                    color = UP if v.startswith("+") else DOWN
                _font(cells[i].paragraphs[0].add_run(str(v)), 9.5, None, color)
                if ri % 2:
                    _shade(cells[i], "F3F5F8")
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph()

    def image(self, path: Path | None, width_cm: float = 16):
        if path and Path(path).exists():
            self.d.add_picture(str(path), width=Cm(width_cm))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def page_break(self):
        self.d.add_page_break()

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(str(path))


FACTOR_ORDER = ["US_TECH", "US_MARKET", "US_RATES", "USD_TWD", "OIL", "CHINA", "JAPAN", "RISK", "GOLD"]


def tuner_section(d: Doc, tune: dict, alpha_out: dict | None = None) -> None:
    d.h1("模型健康度與調整 Agent")
    st = tune["stats"]
    d.table(["指標", "樣本數", "實際表現"], [
        ["大盤隔夜跳空方向命中率", st["taiex_gap_n"], f"{st['taiex_gap_hit'] * 100:.1f}%" if st["taiex_gap_hit"] is not None else "累積中"],
        ["個股隔夜跳空方向命中率", st["stock_gap_n"], f"{st['stock_gap_hit'] * 100:.1f}%" if st["stock_gap_hit"] is not None else "累積中"],
        ["一個月推薦：實際平均超額報酬", st["alpha_n"], f"{st['alpha_excess'] * 100:+.2f}%" if st["alpha_excess"] is not None else "需滿 20 交易日"],
        ["一個月推薦：實際創新高比例", st["alpha_n"], f"{st['alpha_newhigh'] * 100:.1f}%" if st["alpha_newhigh"] is not None else "需滿 20 交易日"],
    ], [7, 2, 4])
    if alpha_out and alpha_out.get("metrics", {}).get("folds"):
        m = alpha_out["metrics"]
        d.para(f"一個月模型滾動回測（{len(m['folds'])} 段時間外樣本）：AUC {m['auc']:.3f}；模型前 20 名一個月內創新高比例 "
               f"{m['topk_newhigh'] * 100:.1f}%（全市場基準 {m['base_rate'] * 100:.1f}%）；前 20 名平均超額報酬 "
               f"{m['topk_excess'] * 100:+.2f}%（平均報酬 {m['topk_ret'] * 100:+.2f}%，同期市場中位數 {m['mkt_ret'] * 100:+.2f}%）。", size=9.5)
        d.table(["測試期間", "AUC", "基準創高率", "前20名創高率", "前20名超額報酬"],
                [[f"{f['start']} ~ {f['end']}", f"{f['auc']:.3f}", f"{f['base_rate'] * 100:.1f}%",
                  f"{f['topk_newhigh'] * 100:.1f}%", f"{f['topk_excess'] * 100:+.2f}%"] for f in m["folds"]], color_col=4)
        d.para("注意：股票池以目前流動性篩選，回測存在存活者偏誤，實際表現可能低於回測。", size=9)
    if tune["issues"]:
        d.bullets([f"[{i['severity']}] {i['detail']}" for i in tune["issues"]], "本次偵測到的失誤")
    if tune["changes"]:
        d.bullets([f"{c['param']}：{c.get('before')} → {c.get('after')}（{c.get('reason', '')}；"
                   f"指標 {c.get('metric_before', '—')} → {c.get('metric_after', '—')}）" for c in tune["changes"]], "本次修正")
    if tune.get("note"):
        d.para("檢討：" + tune["note"])
    if tune.get("recent_log"):
        d.bullets([f"{e['date']}：" + "；".join(c["param"] for c in e["changes"]) for e in tune["recent_log"]], "近期調整紀錄")
    if not tune["issues"] and not tune["changes"]:
        d.para("目前未偵測到重大失誤，模型參數維持不變。")


def build_daily(facts: dict, text: dict, charts: dict, path: Path) -> Path:
    d = Doc()
    d.title(f"台股投資競賽　{facts['title']}", f"{facts['date']}　產出時間 {facts['generated_at']}　"
            f"文字來源：{'LLM（已通過數字核對）' if text['source'] == 'llm-verified' else '程式模板'}")
    pf = facts["portfolio"]

    d.h1("一、今日重點")
    d.para(text["market_summary"])
    d.para(text["outlook"], bold=True)
    d.table(["總資產", "累積報酬", "目前持股率", "執行後持股率", "累計交易筆數", "持有檔數"],
            [[f"{pf['nav']:,.0f}", f"{pf['return_pct']:+.2f}%", f"{pf['holding_ratio'] * 100:.1f}%",
              f"{facts['projected_ratio'] * 100:.1f}%", pf["trade_count"], len(pf["positions"])]], color_col=1)

    # 攻守配置
    al = facts["allocation"]
    d.h1(f"攻守配置：進攻籃 {al['offense'] * 100:.0f}%／穩健籃 {al['defense'] * 100:.0f}%")
    d.para(text.get("allocation_analysis", ""))
    d.bullets(al["reasons"], "比例如何決定")
    names = {"offense": "進攻籃（追強勢、一個月創新高）", "defense": "穩健籃（收正報酬機率高、回檔小、不追高）"}
    d.table(["籃子", "目標金額", "已持有", "本次新買", "執行後佔總資產"],
            [[names[b], f"{v['target']:,}", f"{v['held']:,}", f"{v['new']:,}", f"{v['after_pct']:.1f}%"]
             for b, v in al["baskets"].items()], [5.5, 2.6, 2.4, 2.4, 2.6])
    d.para(f"大盤狀態：加權指數 RSI14 {al['taiex']['rsi14']}、20 日漲跌 {al['taiex']['r20_pct']:+.2f}%、"
           f"距季線 {al['taiex']['dist_ma60_pct']:+.2f}%", size=9)
    if facts.get("def_picks"):
        d.h2("穩健籃候選排行（收正報酬機率高、預期回檔小、未過熱）")
        d.table(["標的", "產業", "收正報酬機率", "預期最大回檔", "β", "日波動", "穩健分數", "訊號"],
                [[f"{p['code']} {p['name']}", p["sector"], f"{p['p_up'] * 100:.1f}%", f"{p['exp_dd'] * 100:.1f}%",
                  f"{p['beta']:.2f}" if p["beta"] is not None else "-",
                  f"{p['vol20'] * 100:.2f}%" if p["vol20"] is not None else "-", f"{p['def_score']:+.2f}", p["signal"]]
                 for p in facts["def_picks"]])

    # 庫存怎麼處理
    adv = facts.get("holding_advice") or []
    if adv:
        d.h1("庫存怎麼處理（加碼／續抱／減碼／全部出清）")
        d.table(["標的", "籃子", "持股", "成本", "現價", "損益", "建議", "數量", "限價"],
                [[f"{h['code']} {h['name']}", h["basket_name"], f"{h['shares']:,}", h["avg_cost"], h["price"],
                  f"{h['pnl_pct']:+.2f}%", h["verdict"], h["qty_text"], h["limit_price"] or "—"] for h in adv],
                [2.6, 1.4, 1.5, 1.4, 1.4, 1.5, 1.5, 4.4, 1.3], color_col=5)
        for h in adv:
            d.h2(f"{h['code']} {h['name']}：{h['verdict']}｜{h['qty_text']}")
            d.bullets(h["reasons"], "理由")
            d.bullets(h["conditions"], "接下來的條件")

    # 市場警示與持股風險
    al_ = facts.get("alerts") or {}
    if al_:
        d.h1("市場警示與持股風險")
        if not al_.get("alerts"):
            d.para("目前沒有需要處理的警示。")
        for a in al_.get("alerts", []):
            color = UP if a["level"] == "嚴重" else (RGBColor(0xA8, 0x6A, 0x00) if a["level"] == "警告" else None)
            d.para(f"［{a['level']}｜{a['scope']}］{a['title']}", bold=True, color=color)
            d.bullets(a["reasons"] + [f"建議：{a['action']}"] + ([a["why_qty"]] if a.get("why_qty") else []))
            for n in (a.get("news") or [])[:2]:
                d.link_bullet(n["title"], n["url"])
        d.para(f"全市場：漲停 {al_.get('limit_up_count', 0)} 檔、跌停 {al_.get('limit_down_count', 0)} 檔"
               + (f"（資料日 {al_['quotes_date']}）" if al_.get("quotes_date") else ""), size=9.5)
        for side, label in (("limit_up", "漲停"), ("limit_down", "跌停")):
            rows = [q for q in al_.get(side, []) if q.get("reason")][:8]
            if rows:
                d.table(["標的", "產業", "漲跌", label + "可能原因"],
                        [[f"{q['code']} {q['name']}", q.get("industry") or "", f"{q['chg_pct']:+.2f}%", q["reason"][:90]]
                         for q in rows], [3, 2.4, 1.6, 9], color_col=2)

    # 二、漲跌歸因
    att = facts["attribution"]
    d.h1(f"二、為何今日會漲／會跌（{att['date']}）")
    if facts["hypothetical"]:
        d.para("目前尚無持股，以下以模型前 5 名候選股示範歸因（金額以 0 股計）。", size=9)
    d.para(text["daily_move_analysis"])
    d.image(charts.get("attribution"))
    if att["rows"]:
        d.table(["標的", "漲跌", "損益(元)", "大盤", "產業", "個股", "主因"],
                [[f"{r['code']} {r['name']}", f"{r['ret_pct']:+.2f}%", f"{r['pnl']:+,}", f"{r['mkt_pct']:+.2f}%",
                  f"{r['ind_pct']:+.2f}%", f"{r['idio_pct']:+.2f}%", r["main_driver"]] for r in att["rows"]], color_col=1)
    for r in att["rows"]:
        d.h2(f"{r['code']} {r['name']}：{r['verdict']}")
        d.bullets(r["explanations"])
    d.image(charts.get("sectors"), 14)

    # 三、虧損補回計畫
    rec = facts["recovery"]
    d.h1("三、虧損如何補回（續抱／加碼／換股／停損）")
    d.para(text["recovery_plan"])
    d.bullets(rec["summary"], "投資組合層級")
    for it in rec["items"]:
        d.h2(f"{it['code']} {it['name']}｜{it['decision']}（未實現 {it['unrealized']:+,} 元，{it['pnl_pct']:+.2f}%）")
        d.bullets(it["lines"])
    if not rec["items"]:
        d.para("目前無持股，無需補回；新部位的停損與換股規則見第六節出場計畫。")

    # 四、國際大事
    d.h1("四、國際大事對持股的影響")
    d.para(text["macro_analysis"])
    if facts["hypothetical"]:
        d.para("（尚無持股：以模型前 5 名候選、每檔假設佔總資產 10% 試算）", size=9)
    for e in facts["macro"]:
        d.h2(f"{e['event']}（{e['horizon']}）")
        d.para(f"對應國際因子：{e['factor_name']}{'上升' if e['direction'] > 0 else '下降'}；"
               f"若因子變動 1 個標準差，持股合計預估影響 {e['total_amount']:+,} 元（總資產 {e['total_pct_nav']:+.2f}%）", bold=True)
        if e["rows"]:
            d.table(["標的", "敏感度β", "相關係數", "預估影響", "金額(元)", "後續決策"],
                    [[f"{r['code']} {r['name']}", r["beta"], r["corr"], f"{r['impact_pct']:+.2f}%", f"{r['amount']:+,}",
                      r["decision"]] for r in e["rows"]], [3, 1.6, 1.6, 1.8, 2, 6], color_col=3)
        for n in e["news"][:3]:
            d.link_bullet(n["title"], n["url"])
    if not facts["macro"]:
        d.para("今日新聞中未辨識出重大國際事件（或未啟用 LLM）。")
    sens = facts.get("macro_sens") or {}
    if sens:
        d.h2("各持股對國際因子的敏感度（因子變動 1 個標準差時的預估漲跌）")
        names = settings()["macro_factors"]
        keys = [k for k in FACTOR_ORDER if any(k in v for v in sens.values())]
        d.table(["標的"] + [names[k]["name"] for k in keys],
                [[f"{c} {facts['names'].get(c, '')}"] + [f"{v[k]['impact_1sd_pct']:+.2f}%" if k in v else "-" for k in keys]
                 for c, v in sens.items()])

    # 五、交易建議
    d.page_break()
    d.h1("五、交易建議總表")
    rows = []
    for o in facts["orders"]:
        qty = f"{o['lots']} 張" + (f" + {o['odd']} 股" if o["odd"] else "")
        rows.append(["買進" if o["action"] == "BUY" else "賣出", o.get("basket_name") or ("穩健籃" if o.get("basket") == "defense" else "進攻籃"),
                     f"{o['code']} {o['name']}", qty, o["limit_price"], o.get("stop", "-"), f"{o['amount']:,}",
                     ("融資" if o.get("margin") else "現股") if o["action"] == "BUY" else "-"])
    if rows:
        d.table(["動作", "籃子", "標的", "數量", "限價", "停損", "金額", "方式"], rows, [1.2, 1.5, 3.2, 2.4, 1.6, 1.6, 2.2, 1.2])
    else:
        d.para("今日無交易建議（無標的達進出場條件）。")
    if facts["funding"]["transfer_from_futures"] > 0:
        d.para(f"資金調度：建議由期貨戶轉入證券戶 {facts['funding']['transfer_from_futures']:,} 元"
               f"（期貨戶保留 {facts['funding']['futures_reserve']:,} 元），以提高持股率。", bold=True)
    d.image(charts.get("allocation"))

    d.h1("六、個股決策說明")
    notes = {n["code"]: n["note"] for n in text["order_notes"]}
    for o in facts["orders"]:
        if o["action"] == "BUY":
            d.h2(f"買進【{o.get('basket_name', '')}】{o['code']} {o['name']}｜{o['lots']} 張" + (f" + {o['odd']} 股" if o["odd"] else "")
                 + f"｜{'穩健分數' if o.get('basket') == 'defense' else '綜合分數'} {o['score']:+.2f}")
            if o["code"] in notes:
                d.para(notes[o["code"]], color=NAVY)
            d.bullets(o["why_stock"], "為何選擇這檔")
            d.bullets(o["why_qty"], "為何買這個股數（計算過程）")
            if o["why_more"]:
                d.bullets(o["why_more"], "多買的原因")
            if o["why_less"]:
                d.bullets(o["why_less"], "少買的原因")
            if o["risks"]:
                d.bullets(o["risks"], "風險與不利因子")
            d.bullets(o["exit_plan"] + [f"來回交易成本約 {o['round_trip_cost_pct']:.3f}%，損益兩平價 {o['breakeven']} 元"],
                      "出場計畫（何時賣、賣幾張）")
        else:
            d.h2(f"賣出 {o['code']} {o['name']}｜{o['lots']} 張" + (f" + {o['odd']} 股" if o["odd"] else ""))
            d.bullets([o["why_sell"]], "賣出原因")
            d.bullets([o["why_qty"], f"賣出後剩餘 {o['holding_after']:,} 股；預估稅費 {o['est_cost']:,.0f} 元"], "為何賣這個股數")

    # 七、全市場排行
    d.page_break()
    am = facts["alpha_model"]
    d.h1("七、全市場一個月創新高潛力排行")
    d.para(f"掃描 {am['n_stocks']} 檔流動性足夠的上市櫃股票與 ETF（資料截至 {am['asof']}），預測未來 20 個交易日內"
           f"創 52 週新高的機率與超額報酬。全市場歷史基準創新高率 {am['base_rate'] * 100:.1f}%。")
    d.image(charts.get("top_picks"))
    d.para("訊號「過熱不追」：RSI 或 20 日漲幅超過上限，即使模型分數高也不買，避免全部追高。", size=9)
    d.table(["標的", "產業", "創新高機率", "預期超額", "距高點", "營收年增", "綜合分數", "訊號"],
            [[f"{p['code']} {p['name']}", p["sector"], f"{p['p_newhigh'] * 100:.1f}%", f"{p['exp_excess'] * 100:+.2f}%",
              f"{p['dist_hi'] * 100:+.1f}%", f"{p['rev_yoy'] * 100:+.1f}%" if p["rev_yoy"] is not None else "-",
              f"{p['score']:+.2f}", p["signal"]] for p in facts["top_picks"]], color_col=3)

    # 八、美股與隔夜預測
    d.h1("八、美股隔夜與開盤預測")
    d.image(charts.get("us"))
    us_rows = [[r["name"], r["ticker"], r["close"], f"{r['change_pct']:+.2f}%",
                f"{r['chg_5d_pct']:+.2f}%" if r["chg_5d_pct"] is not None else "-"] for r in facts["us_snapshot"]]
    d.table(["名稱", "代號", "收盤", "日漲跌", "5日漲跌"], us_rows, color_col=3)
    d.image(charts.get("pred"))
    d.para("方法：Ridge 迴歸，特徵為前一美股交易日的費半、那斯達克期貨、台積電 ADR、VIX、匯率、產業對應美股與前日自身報酬；"
           "命中率為樣本外方向命中率，需高於 55% 且高於「永遠猜多數方向」的基準才採用。", size=9)

    # 九、新聞
    d.h1("九、個股新聞事件")
    if facts["events"]:
        for e in facts["events"]:
            tag = "＋" * max(e["sentiment"], 0) + "－" * max(-e["sentiment"], 0) or "＝"
            d.para(f"{tag} [{', '.join(e['codes'])}] {e['event']}" + ("（多來源）" if e["multi_source"] else ""), bold=True)
            for n in e["news"][:3]:
                d.link_bullet(n["title"], n["url"])
    else:
        for n in facts["news"][:15]:
            d.link_bullet(f"({n['source']}) {n['title']}", n["url"])
    if facts.get("filings"):
        d.h1("法說會 / 財報重點（PDF）")
        for code, docs in facts["filings"].items():
            for f in docs:
                d.h2(f"{code} {f['company']} {f['period']}（{f['file']}）")
                d.bullets([f"{p['point']}（第 {p['page']} 頁）" for p in f["points"]])

    # 十、組合與規則、模型健康
    d.page_break()
    d.h1("十、投資組合與競賽規則檢查")
    if pf["positions"]:
        d.table(["標的", "籃子", "股數", "成本", "現價", "損益", "權重"],
                [[f"{p['code']} {p['name']}", p.get("basket", ""), f"{p['shares']:,}", p["avg_cost"], p["price"],
                  f"{p['pnl_pct']:+.2f}%", f"{p['weight_pct']:.1f}%"] for p in pf["positions"]], color_col=5)
    d.table(["檢查項目", "結果", "說明"], [[c["item"], c["status"], c["detail"]] for c in facts["compliance"]], [4.5, 1.6, 10])
    d.image(charts.get("nav"))
    tuner_section(d, facts["tuner"], facts["alpha_model"])

    d.h1("附錄：方法與聲明")
    d.bullets([
        "所有價格、張數、金額、比例、敏感度皆由程式依公開資料計算；LLM 只負責新聞事件辨識與文字撰寫，並經程式逐一核對數字。",
        "資料來源：Yahoo Finance（美股、全市場日 K）、FinMind（股票清單、三大法人、月營收）、臺灣證券交易所／櫃買中心、"
        "鉅亨網、中央社、Yahoo 股市、Google 新聞。",
        "本報告為模擬投資競賽之研究用途，不構成投資建議。",
    ])
    d.save(path)
    return path
